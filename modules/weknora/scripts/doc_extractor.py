#!/usr/bin/env python3
"""
Local Standalone Document Extractor for WeKnora Lite.
Extracts clean Markdown / text from PDF, DOCX, XLSX, PPTX, and EPUB files
without needing Docker or any external cloud services.
"""

import sys
import os
import io
import re
import shutil
import subprocess
import tempfile
import unicodedata
import base64
import json
import urllib.request
import urllib.error

_CACHED_VISION_MODEL = None
_VISION_CHECK_DONE = False

def get_available_vision_model() -> str:
    """Checks local Ollama server for available vision models."""
    global _CACHED_VISION_MODEL, _VISION_CHECK_DONE
    if _VISION_CHECK_DONE:
        return _CACHED_VISION_MODEL or ""
    _VISION_CHECK_DONE = True
    try:
        req = urllib.request.Request("http://127.0.0.1:11434/api/tags", headers={"User-Agent": "doc_extractor"})
        with urllib.request.urlopen(req, timeout=2) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            models = [m.get("name", "") for m in data.get("models", [])]
            for candidate in ["minicpm-v:latest", "minicpm-v", "gemma4:e4b", "llava:latest", "llava:7b", "glm-ocr:latest"]:
                for m in models:
                    if m == candidate or m.startswith(candidate.split(":")[0]):
                        _CACHED_VISION_MODEL = m
                        return _CACHED_VISION_MODEL
    except Exception:
        pass
    return ""

def describe_image_with_vlm(img_bytes: bytes, caption_hint: str = "") -> str:
    """Calls local Ollama VLM to generate a concise technical description of a diagram/chart/image."""
    model = get_available_vision_model()
    if not model or not img_bytes:
        return ""
    try:
        b64 = base64.b64encode(img_bytes).decode("utf-8")
        prompt = (
            "Describe this figure or diagram concisely in 2-3 sentences for a technical search index. "
            "Explicitly mention the figure title or number if visible, axis names and units, key curves, labels, or flow elements."
        )
        if caption_hint:
            prompt += f"\nCaption context: {caption_hint}"
        payload = {
            "model": model,
            "prompt": prompt,
            "images": [b64],
            "stream": False
        }
        req = urllib.request.Request(
            "http://127.0.0.1:11434/api/generate",
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"}
        )
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            return data.get("response", "").strip()
    except Exception:
        return ""

def find_tesseract_binary() -> str:
    """Finds local Tesseract OCR executable if available."""
    env_cmd = os.environ.get("TESSERACT_CMD")
    if env_cmd and os.path.isfile(env_cmd):
        return env_cmd
    which_cmd = shutil.which("tesseract") or shutil.which("tesseract.exe")
    if which_cmd and os.path.isfile(which_cmd):
        return which_cmd
    candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    home = os.environ.get("USERPROFILE") or os.environ.get("HOME")
    if home:
        candidates.append(os.path.join(home, "AppData", "Local", "Programs", "Tesseract-OCR", "tesseract.exe"))
    for c in candidates:
        if os.path.isfile(c):
            return c
    return ""

def try_tesseract_ocr(page, lang="eng", dpi=200) -> str:
    """Renders page to image and performs OCR if Tesseract is available."""
    tess_bin = find_tesseract_binary()
    if not tess_bin:
        return ""
    try:
        pix = page.get_pixmap(dpi=dpi)
        img_bytes = pix.tobytes("png")
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_in:
            tmp_path = tmp_in.name
            tmp_in.write(img_bytes)
        try:
            cmd = [tess_bin, tmp_path, "stdout", "-l", lang]
            proc = subprocess.run(cmd, capture_output=True, timeout=60, check=False)
            out = proc.stdout.decode("utf-8", errors="replace").strip()
            return out
        finally:
            if os.path.isfile(tmp_path):
                try:
                    os.unlink(tmp_path)
                except Exception:
                    pass
    except Exception:
        return ""


def normalize_math_unicode(text: str) -> str:
    """Normalizes Mathematical Alphanumeric Symbols (U+1D400 to U+1D7FF) to standard ASCII equivalents."""
    result = []
    for ch in text:
        code = ord(ch)
        if 0x1D400 <= code <= 0x1D419 or 0x1D434 <= code <= 0x1D44D:
            offset = 0x1D434 if code >= 0x1D434 else 0x1D400
            result.append(chr(ord('A') + (code - offset)))
        elif 0x1D41A <= code <= 0x1D433 or 0x1D44E <= code <= 0x1D467:
            offset = 0x1D44E if code >= 0x1D44E else 0x1D41A
            result.append(chr(ord('a') + (code - offset)))
        elif ch in ['\u2212', '−']:
            result.append('-')
        elif ch in ['\u2217', '∗', '×']:
            result.append('*')
        else:
            decomp = unicodedata.normalize('NFKD', ch)
            result.append(decomp if decomp else ch)
    return "".join(result)

def extract_pdf(path: str) -> str:
    import pymupdf
    doc = pymupdf.open(path)
    pages = []

    for i, page in enumerate(doc):
        # Normalize page rotation if landscape/rotated (e.g. standards with Rotate=270/90)
        orig_rot = page.rotation
        if orig_rot != 0:
            tabs_orig = page.find_tables()
            page.set_rotation(0)
            tabs_0 = page.find_tables()
            # If rotation 0 yields valid tables or cleaner orientation, keep 0; else restore
            if len(tabs_0.tables) >= len(tabs_orig.tables):
                tabs = tabs_0
            else:
                page.set_rotation(orig_rot)
                tabs = tabs_orig
        else:
            tabs = page.find_tables()

        table_bboxes = []
        table_dict = {}

        if tabs.tables:
            for t in tabs.tables:
                tb = tuple(t.bbox)
                table_bboxes.append(tb)
                table_dict[tb] = t.to_markdown()

        # 2. Extract drawings for radicals and horizontal rules
        drawings = page.get_drawings()
        sqrt_rects = []
        for dr in drawings:
            r = dr.get("rect")
            items = dr.get("items", [])
            if not r:
                continue
            w = r[2] - r[0]
            h = r[3] - r[1]
            if len(items) >= 2 and w > 15 and h > 8:
                sqrt_rects.append(r)

        # 2b. Extract vector drawing fraction lines, vertical modulus bars, and assemble LaTeX equations outside tables
        words = page.get_text("words")
        frac_lines = []
        vert_bars = []
        for d in drawings:
            r = d.get("rect")
            if not r:
                continue
            if any(r.x0 >= tb[0] - 5 and r.x1 <= tb[2] + 5 and r.y0 >= tb[1] - 5 and r.y1 <= tb[3] + 5 for tb in table_bboxes):
                continue
            if r.y0 < 110 or r.y1 > 750:
                continue
            # Fraction line
            if r.height <= 3.0 and r.width >= 8.0:
                frac_lines.append(r)
            # Vertical stroke bar (modulus / absolute value / norm)
            elif r.width <= 2.5 and r.height >= 8.0:
                vert_bars.append(r)

        # Group into equation rows
        frac_rows = {}
        for fl in frac_lines:
            y_key = round(fl.y0 / 12.0) * 12.0
            frac_rows.setdefault(y_key, []).append(fl)

        combined_items = []
        equation_suppress_rects = []

        # Add tables to page items
        for tb, md in table_dict.items():
            if md.strip():
                combined_items.append((tb[1], "table", md))

        for y_key, lines in frac_rows.items():
            avg_y = sum((l.y0 + l.y1) / 2.0 for l in lines) / len(lines)
            row_words = [w for w in words if abs((w[1] + w[3]) / 2.0 - avg_y) <= 22.0]
            if not row_words:
                continue

            frac_map = {}
            used_words = set()
            for fl in lines:
                above = []
                below = []
                for w in row_words:
                    cx = (w[0] + w[2]) / 2.0
                    cy = (w[1] + w[3]) / 2.0
                    if (fl.x0 - 4.0 <= cx <= fl.x1 + 4.0 or
                        (w[0] >= fl.x0 - 2.0 and w[2] <= fl.x1 + 2.0) or
                        (w[2] > fl.x0 and w[0] < fl.x1 and (w[2] - max(w[0], fl.x0)) >= 0.5 * (w[2] - w[0]))):
                        txt = normalize_math_unicode(w[4])
                        if cy < avg_y:
                            above.append((w[0], txt))
                        elif cy > avg_y:
                            below.append((w[0], txt))

                # Include vertical modulus bars (|) within horizontal span of fraction
                for vb in vert_bars:
                    if fl.x0 - 4.0 <= vb.x0 <= fl.x1 + 4.0:
                        if vb.y1 <= avg_y + 3.0 and abs(vb.y1 - avg_y) <= 25.0:
                            above.append((vb.x0, "|"))
                        elif vb.y0 >= avg_y - 3.0 and abs(vb.y0 - avg_y) <= 25.0:
                            below.append((vb.x0, "|"))

                if not above or not below:
                    continue

                above.sort(key=lambda x: x[0])
                below.sort(key=lambda x: x[0])

                num_str = " ".join(x[1] for x in above).strip()
                den_str = " ".join(x[1] for x in below).strip()

                # Clean up spacing around | and parentheses and variable subscripts
                for target_str in ("num_str", "den_str"):
                    val = locals()[target_str]
                    val = re.sub(r'\|\s*', '|', val)
                    val = re.sub(r'\s*\|', '|', val)
                    val = re.sub(r'\(\s*', '(', val)
                    val = re.sub(r'\s*\)', ')', val)
                    val = re.sub(r'\s*−\s*', ' - ', val)
                    val = re.sub(r'\s*\+\s*', ' + ', val)
                    val = re.sub(r'f\s*([0-9])', r'f_{\1}', val)
                    val = re.sub(r'p\s*\(', r'p(', val)
                    if target_str == "num_str":
                        num_str = val
                    else:
                        den_str = val

                if not num_str or not den_str:
                    continue

                prefix = ""
                if num_str.startswith("(") and not num_str.endswith(")"):
                    prefix = "("
                    num_str = num_str[1:].strip()

                frac_latex = f"{prefix}\\frac{{{num_str}}}{{{den_str}}}"
                frac_map[fl.x0] = frac_latex
                for w in row_words:
                    cx = (w[0] + w[2]) / 2.0
                    if fl.x0 - 4.0 <= cx <= fl.x1 + 4.0:
                        used_words.add(id(w))

            if not frac_map:
                continue

            elements = []
            for fl in lines:
                if fl.x0 in frac_map:
                    elements.append((fl.x0, frac_map[fl.x0]))
            for w in row_words:
                if id(w) not in used_words:
                    cy = (w[1] + w[3]) / 2.0
                    if abs(cy - avg_y) <= 8.0:
                        txt = normalize_math_unicode(w[4])
                        elements.append((w[0], txt))

            elements.sort(key=lambda x: x[0])
            eq_text = " ".join(e[1] for e in elements)
            if eq_text:
                # Format left side subscripts (e.g. D DF 2- -> D_{DF2-})
                eq_text = re.sub(r'D\s+DF\s+([0-9][\+\-]*)', r'D_{DF\1}', eq_text)
                combined_items.append((avg_y, "formula", f"$$\n{eq_text}\n$$"))
                min_x = min(w[0] for w in row_words)
                max_x = max(w[2] for w in row_words)
                min_y = min(w[1] for w in row_words)
                max_y = max(w[3] for w in row_words)
                equation_suppress_rects.append((min_x - 5, min_y - 5, max_x + 5, max_y + 5))

        # 3. Extract text blocks
        d = page.get_text("dict")
        blocks = d.get("blocks", [])

        # 3b. Extract figures and diagrams with local VLM
        figure_suppress_rects = []
        fig_caption_pattern = re.compile(r"^(?:Figure|Fig\.|Chart|Diagram)\s+([0-9]+(?:\.[0-9]+)*)", re.IGNORECASE)
        vlm_model = get_available_vision_model()
        fig_count_on_page = 0

        def _calc_dr_bbox(dr):
            pts = []
            r = dr.get("rect")
            if r and (r.width > 0 or r.height > 0):
                pts.extend([pymupdf.Point(r.x0, r.y0), pymupdf.Point(r.x1, r.y1)])
            for it in dr.get("items", []):
                for elem in it[1:]:
                    if isinstance(elem, pymupdf.Point):
                        pts.append(elem)
                    elif isinstance(elem, pymupdf.Rect):
                        pts.extend([pymupdf.Point(elem.x0, elem.y0), pymupdf.Point(elem.x1, elem.y1)])
            if pts:
                return pymupdf.Rect(
                    min(p.x for p in pts),
                    min(p.y for p in pts),
                    max(p.x for p in pts),
                    max(p.y for p in pts)
                )
            return pymupdf.Rect(r) if r else pymupdf.Rect(0, 0, 0, 0)

        drawing_rects = [_calc_dr_bbox(dr) for dr in drawings] if drawings else []

        if vlm_model:
            # First pass: find explicit figure captions (e.g. Figure 1.1 ...)
            for b in blocks:
                if b.get("type") != 0 or fig_count_on_page >= 4:
                    continue
                raw_text = "".join(s.get("text", "") for l in b.get("lines", []) for s in l.get("spans", [])).strip()
                if fig_caption_pattern.match(raw_text):
                    cap_rect = pymupdf.Rect(b["bbox"])
                    # Look for vector drawings or raster images sitting above the caption (or below)
                    above_cand = [
                        r for r in drawing_rects
                        if r.y1 <= cap_rect.y0 + 15 and r.y0 >= cap_rect.y0 - 450
                        and max(r.width, r.height) >= 15
                    ] + [
                        pymupdf.Rect(ib["bbox"]) for ib in blocks
                        if ib.get("type") == 1 and pymupdf.Rect(ib["bbox"]).y1 <= cap_rect.y0 + 15
                        and pymupdf.Rect(ib["bbox"]).y0 >= cap_rect.y0 - 450
                    ]
                    below_cand = [
                        r for r in drawing_rects
                        if r.y0 >= cap_rect.y1 - 15 and r.y1 <= cap_rect.y1 + 450
                        and max(r.width, r.height) >= 15
                    ] + [
                        pymupdf.Rect(ib["bbox"]) for ib in blocks
                        if ib.get("type") == 1 and pymupdf.Rect(ib["bbox"]).y0 >= cap_rect.y1 - 15
                        and pymupdf.Rect(ib["bbox"]).y1 <= cap_rect.y1 + 450
                    ]

                    cand_rects = above_cand if above_cand else below_cand
                    if cand_rects:
                        u_rect = cand_rects[0]
                        for cr in cand_rects[1:]:
                            u_rect |= cr
                        crop_rect = pymupdf.Rect(
                            max(0, min(u_rect.x0, cap_rect.x0) - 30),
                            max(0, min(u_rect.y0, cap_rect.y0) - 15),
                            min(page.rect.width, max(u_rect.x1, cap_rect.x1) + 30),
                            min(page.rect.height, max(u_rect.y1, cap_rect.y1) + 15)
                        )
                        pix = page.get_pixmap(clip=crop_rect, dpi=120)
                        vlm_desc = describe_image_with_vlm(pix.tobytes("png"), caption_hint=raw_text)
                        if vlm_desc:
                            fig_md = f"> **{raw_text}**\n>\n> *Visual Analysis*: {vlm_desc}"
                            combined_items.append((crop_rect.y0, "figure", fig_md))
                            figure_suppress_rects.append(crop_rect)
                            fig_count_on_page += 1

            # Second pass: stand-alone raster image blocks (type == 1) not covered by captioned figures
            for b in blocks:
                if b.get("type") != 1 or fig_count_on_page >= 4:
                    continue
                ib_rect = pymupdf.Rect(b.get("bbox", (0, 0, 0, 0)))
                if ib_rect.width >= 100 and ib_rect.height >= 80:
                    if any(ib_rect.intersects(sr) for sr in figure_suppress_rects):
                        continue
                    pix = page.get_pixmap(clip=ib_rect, dpi=120)
                    vlm_desc = describe_image_with_vlm(pix.tobytes("png"))
                    if vlm_desc:
                        img_md = f"> **[Diagram / Figure]**\n>\n> *Visual Analysis*: {vlm_desc}"
                        combined_items.append((ib_rect.y0, "figure", img_md))
                        figure_suppress_rects.append(ib_rect)
                        fig_count_on_page += 1

        for b in blocks:
            if b.get("type") != 0:
                continue
            lines = b.get("lines", [])
            if not lines:
                continue

            b_bbox = b.get("bbox", (0, 0, 0, 0))

            # Exclude text that falls inside a recognized figure region (captions and internal axis/curve labels)
            if any(b_bbox[1] >= sr.y0 - 5 and b_bbox[3] <= sr.y1 + 5 and max(b_bbox[0], sr.x0) < min(b_bbox[2], sr.x1) for sr in figure_suppress_rects):
                continue

            # Exclude text that is already captured inside a table bounding box
            if any(b_bbox[0] >= tb[0] - 5 and b_bbox[2] <= tb[2] + 5 and b_bbox[1] >= tb[1] - 5 and b_bbox[3] <= tb[3] + 5 for tb in table_bboxes):
                continue

            # Exclude text that is already captured inside an equation row (check vertical & horizontal overlap)
            if any(abs((b_bbox[1] + b_bbox[3]) / 2.0 - ((sr[1] + sr[3]) / 2.0)) <= 15.0 and max(b_bbox[0], sr[0]) < min(b_bbox[2], sr[2]) for sr in equation_suppress_rects):
                continue

            bh = b_bbox[3] - b_bbox[1]
            raw_block_text = "".join(s.get("text", "") for l in lines for s in l.get("spans", []))

            # Single-row math formula with sub/superscripts or radical
            is_formula = (bh <= 35 and len(lines) >= 2 and ("=" in raw_block_text or "THD" in raw_block_text or any(
                abs(r[1] - b_bbox[1]) < 15 and r[0] >= b_bbox[0] - 10 and r[2] <= b_bbox[2] + 15 for r in sqrt_rects
            )))

            if is_formula:
                spans = [s for l in lines for s in l.get("spans", []) if s.get("text")]
                if spans:
                    spans.sort(key=lambda s: s["bbox"][0])
                    base_size = max(s["size"] for s in spans)

                    has_sqrt = any(
                        abs(r[1] - b_bbox[1]) < 15 and r[0] >= b_bbox[0] - 10 and r[2] <= b_bbox[2] + 15
                        for r in sqrt_rects
                    )

                    tokens = []
                    for s in spans:
                        txt = s["text"]
                        sz = s["size"]
                        y0 = s["bbox"][1]
                        if sz < base_size * 0.82:
                            if y0 > b_bbox[1] + bh * 0.45:
                                tokens.append(f"_{txt}")
                            else:
                                tokens.append(f"^{txt}")
                        else:
                            tokens.append(txt)

                    assembled = "".join(tokens).strip()
                    assembled = assembled.replace("[ %]", "[%]").replace("( ", "(").replace(" )", ")")

                    if has_sqrt and "=" in assembled:
                        parts = assembled.split("=", 1)
                        left = parts[0].strip()
                        right = parts[1].strip()
                        if right.startswith("100"):
                            expr = right[3:].strip()
                            if "/" in expr:
                                num, den = expr.split("/", 1)
                                num = num.strip().lstrip("(").rstrip(")")
                                den = den.strip()
                                formula_latex = f"{left} = 100 \\sqrt{{\\frac{{{num}}}{{{den}}}}}"
                            else:
                                formula_latex = f"{left} = 100 \\sqrt{{{expr}}}"
                        else:
                            if "/" in right:
                                num, den = right.split("/", 1)
                                num = num.strip().lstrip("(").rstrip(")")
                                den = den.strip()
                                formula_latex = f"{left} = \\sqrt{{\\frac{{{num}}}{{{den}}}}}"
                            else:
                                formula_latex = f"{left} = \\sqrt{{{right}}}"
                        assembled = f"$$\n{formula_latex}\n$$"
                    elif "=" in assembled and ("^" in assembled or "_" in assembled):
                        assembled = f"$$\n{assembled}\n$$"

                    combined_items.append((b_bbox[1], "formula", assembled))
            else:
                # Group multi-line text by vertical baseline y-coordinate (e.g. definition lists)
                rows = {}
                for l in lines:
                    y = round(l["bbox"][1], 1)
                    found = False
                    for r_y in rows:
                        if abs(r_y - y) < 4:
                            rows[r_y].extend(l.get("spans", []))
                            found = True
                            break
                    if not found:
                        rows[y] = list(l.get("spans", []))

                sorted_rows = sorted(rows.items(), key=lambda x: x[0])
                block_text_lines = []
                for r_y, row_spans in sorted_rows:
                    row_spans.sort(key=lambda s: s["bbox"][0])
                    row_txt = "".join(normalize_math_unicode(s["text"]) for s in row_spans).strip()
                    if row_txt:
                        block_text_lines.append(row_txt)
                if block_text_lines:
                    combined_items.append((b_bbox[1], "text", "\n".join(block_text_lines)))

        # Sort all items (tables and text) by vertical y0 coordinate
        combined_items.sort(key=lambda x: x[0])
        page_text = "\n\n".join(item[2] for item in combined_items).strip()
        if not page_text or len(page_text.split()) < 8:
            ocr_res = try_tesseract_ocr(page)
            if ocr_res:
                page_text = f"{page_text}\n\n{ocr_res}".strip() if page_text else ocr_res

        if page_text:
            pages.append(f"<!-- Page {i+1} -->\n{page_text}")

    doc.close()
    return "\n\n".join(pages)

def extract_docx(path: str) -> str:
    import docx
    doc = docx.Document(path)
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style_name = getattr(p.style, "name", "")
        if style_name.startswith("Heading 1"):
            lines.append(f"# {text}")
        elif style_name.startswith("Heading 2"):
            lines.append(f"## {text}")
        elif style_name.startswith("Heading 3"):
            lines.append(f"### {text}")
        else:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_vals = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(row_vals):
                lines.append(" | ".join(row_vals))
    return "\n\n".join(lines)

def extract_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for name in wb.sheetnames:
        ws = wb[name]
        lines.append(f"## Sheet: {name}")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None and str(c).strip() for c in row):
                lines.append(" | ".join(str(c).strip().replace("\n", " ") if c is not None else "" for c in row))
    return "\n\n".join(lines)

def extract_pptx(path: str) -> str:
    import pptx
    prs = pptx.Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides):
        s_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        s_text.append(t)
        if s_text:
            slides.append(f"## Slide {i+1}\n" + "\n".join(s_text))
    return "\n\n".join(slides)

def extract_epub(path: str) -> str:
    import pymupdf
    doc = pymupdf.open(path)
    pages = []
    for page in doc:
        t = page.get_text("text").strip()
        if t:
            pages.append(t)
    return "\n\n".join(pages)

def main():
    if len(sys.argv) < 3:
        print("Usage: python doc_extractor.py <file_path> <file_type>", file=sys.stderr)
        sys.exit(1)

    file_path = sys.argv[1]
    file_type = sys.argv[2].lower().lstrip(".")

    if not os.path.exists(file_path):
        print(f"File not found: {file_path}", file=sys.stderr)
        sys.exit(1)

    try:
        if file_type == "pdf":
            content = extract_pdf(file_path)
        elif file_type in ("docx", "doc"):
            content = extract_docx(file_path)
        elif file_type in ("xlsx", "xls"):
            content = extract_xlsx(file_path)
        elif file_type in ("pptx", "ppt"):
            content = extract_pptx(file_path)
        elif file_type == "epub":
            content = extract_epub(file_path)
        else:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()

        sys.stdout.buffer.write(content.encode("utf-8"))
    except Exception as e:
        print(f"Extraction error for {file_type}: {e}", file=sys.stderr)
        sys.exit(2)

if __name__ == "__main__":
    main()
